from pathlib import Path

from docx import Document

from chunk import chunk_text
from extract import extract_text


def test_extract_text_handles_txt(tmp_path):
    path = tmp_path / "sample.txt"
    path.write_text("This is a plain text file.\nIt contains multiple lines.")

    result = extract_text(str(path))

    assert "plain text file" in result
    assert "multiple lines" in result


def test_extract_text_handles_docx(tmp_path):
    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Quarterly revenue was $20M.")
    doc.add_paragraph("The company expanded into Europe.")
    doc.save(path)

    result = extract_text(str(path))

    assert "Quarterly revenue was $20M." in result
    assert "expanded into Europe" in result


def test_chunk_text_keeps_headings_and_respects_overlap():
    text = "# Revenue\nThe company generated significant revenue in 2024.\n\n# Risk\nSupply chain interruptions remain a key concern."

    chunks = chunk_text(text, chunk_size=80, overlap=20)

    assert len(chunks) > 1
    assert any("Revenue" in chunk for chunk in chunks)
    assert any("Risk" in chunk for chunk in chunks)
    assert all(len(chunk) <= 80 for chunk in chunks)
